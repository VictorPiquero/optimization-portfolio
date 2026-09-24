from read_instance import read_solomon_instance
from gurobi_model import solve_cvrptw
from openpyxl import Workbook, load_workbook
import math

wb = Workbook()
ws = wb.active
ws.title = "Gurobi"

ws.append([
    "customers",
    "K_ref",
    "fleet_factor",
    "vehicles",
    "solver",
    "runtime",
    "objective",
    "best_bound",
    "gap",
    "status"
])

wb.save("results_gurobi.xlsx")

(
    all_nodes,
    all_customers,
    max_vehicles,
    Q,
    all_demand,
    all_earliest_time,
    all_latest_time,
    all_service_time,
    all_cost
) = read_solomon_instance("02_routing_problems\\CVRPTW\\Compare_programs\\instances\\C1_6_1.TXT")

# Feasible reference fleet obtained in preliminary experiments
experiments = {
    25: 7,
    50: 10,
    100: 17,
    200: 31,
    300: 38,
    400: 49,
    500: 56,
}


for n, k_ref in experiments.items():

    customers = all_customers[:n]
    nodes = [0] + customers

    demand = {
        i: all_demand[i]
        for i in customers
    }

    earliest_time = {
        i: all_earliest_time[i]
        for i in customers
    }

    latest_time = {
        i: all_latest_time[i]
        for i in nodes
    }

    service_time = {
        i: all_service_time[i]
        for i in nodes
    }

    cost = {
        (i,j): all_cost[i,j]
        for i in nodes
        for j in nodes
        if i != j
    }


    total_demand = sum(
    demand[i]
    for i in customers
    )


    k_20 = math.ceil(k_ref * 1.20)
    k_50 = math.ceil(k_ref * 1.50)
    max_vehicles = len(customers)

    for fleet_factor, k in [
    ("20%", k_20),
    ("50%", k_50),
    ("100%", max_vehicles)
    ]:

        print(f"Solving CVRPTW with {n} customers and fleet factor {fleet_factor}...")

        M_customers = max(
            latest_time[i] + service_time[i] + cost[i,j] - earliest_time[j]
            for i in customers
            for j in customers
            if i != j
        )

        M_depot = max(
            service_time[0] + cost[0,j] - earliest_time[j]
            for j in customers
        )

        M = max(0, M_customers, M_depot)

        result = solve_cvrptw(
            nodes,
            customers,
            k,
            Q,
            demand,
            earliest_time,
            latest_time,
            service_time,
            cost,
            M,
            time_limit=1800
        )

        wb = load_workbook("results_gurobi.xlsx")
        ws = wb["Gurobi"]

        ws.append([
            n,
            k_ref,
            fleet_factor,
            k,
            "Gurobi",
            result["runtime"],
            result["objective"],
            result["best_bound"],
            result["gap"],
            result["status"]
        ])

        wb.save("results_gurobi.xlsx")
        wb.close()


